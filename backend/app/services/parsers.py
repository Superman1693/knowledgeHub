"""文档解析器：按扩展名分派，输出纯文本（尽量保留结构与页码）。

设计目标：可扩展。新增格式只需注册一个解析函数。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from app.core.exceptions import ParseError
from app.core.logging import get_logger
from app.services import ocr_service
from app.storage import files

logger = get_logger(__name__)

# 纯文本类扩展名：直接按文本读取
TEXT_EXTENSIONS = {
    "md", "markdown", "txt", "py", "js", "ts", "jsx", "tsx",
    "html", "htm", "css", "json", "yaml", "yml", "ini", "toml",
    "conf", "cfg", "env", "java", "go", "rs", "c", "cpp", "h", "hpp",
    "cs", "php", "rb", "sql", "sh", "bat", "xml", "csv", "log",
}


@dataclass
class ParsedDocument:
    """解析结果。"""

    text: str
    # 逐页文本（仅 PDF 等分页格式提供），用于保留页码
    pages: list[str] = field(default_factory=list)
    # 结构提示：如从 Markdown 提取的标题行
    headings: list[str] = field(default_factory=list)


def _parse_text(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    headings: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            headings.append(stripped.lstrip("#").strip())
    return ParsedDocument(text=text, headings=headings)


def _parse_pdf(path: Path) -> ParsedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少 PDF 解析依赖 pypdf，请先安装依赖") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PDF 打开失败：{exc}") from exc

    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - 单页失败不影响整体
            pages.append("")

    # 扫描件回退：整篇无文本时尝试 OCR（仅对确实为空的页面做，避免浪费）
    ocr_pages = 0
    if not any(p.strip() for p in pages):
        ocr_pages = _ocr_pdf_pages(path, pages)

    text = "\n\n".join(pages).strip()
    if not text:
        status = ocr_service.ocr_status()
        hint = status["hint"] or "未提取到文本"
        raise ParseError(f"PDF 未提取到文本内容（{hint}）")
    return ParsedDocument(text=text, pages=pages)


def _parse_image(path: Path) -> ParsedDocument:
    """图片解析：走 OCR 把截图里的文字提取成正文。

    此前 OCR **只作用于 PDF 页面**，截图（PNG/JPG）连作为文档导入都不允许 ——
    而「把报错截图里的文字整进知识库」正是技术笔记最高频的场景。
    这里复用 `ocr_service.ocr_image_bytes()`（PDF 页面识别用的就是它）。
    """
    status = ocr_service.ocr_status()
    if not status["available"]:
        raise ParseError(
            f"图片需要 OCR 才能提取文字，但当前 OCR 不可用（{status['hint']}）"
        )
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise ParseError(f"读取图片失败：{exc}") from exc
    try:
        text = ocr_service.ocr_image_bytes(data)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"图片 OCR 失败：{exc}") from exc

    text = (text or "").strip()
    if not text:
        raise ParseError("图片未识别到文字（可能是纯图形截图，没有可提取的文本）")
    return ParsedDocument(text=text)


def _ocr_pdf_pages(path: Path, pages: list[str]) -> int:
    """对无文本层的 PDF 页面做 OCR，原地写回 pages，返回成功识别的页数。

    只对确实为空的页面调用 OCR；引擎不可用时返回 0，保持原有行为。
    """
    status = ocr_service.ocr_status()
    if not status["available"] or not status["renderer_installed"]:
        logger.warning("PDF 疑似扫描件，但 OCR 不可用：%s", status["hint"])
        return 0

    rendered = ocr_service.render_pdf_pages(str(path))
    if not rendered:
        return 0

    done = 0
    for page_no, png in rendered:
        idx = page_no - 1
        if idx < len(pages) and not pages[idx].strip():
            try:
                recognized = ocr_service.ocr_image_bytes(png)
            except Exception as exc:  # noqa: BLE001
                logger.warning("第 %d 页 OCR 失败：%s", page_no, exc)
                continue
            if recognized:
                pages[idx] = recognized
                done += 1
    if done:
        logger.info("OCR 成功识别 %d/%d 页", done, len(rendered))
    return done


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少 DOCX 解析依赖 python-docx，请先安装依赖") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"DOCX 打开失败：{exc}") from exc

    lines: list[str] = []
    headings: list[str] = []
    for para in document.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading") or style.startswith("标题"):
            # 保留标题结构，用 Markdown 标题标记
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"{'#' * min(int(level), 6)} {txt}")
            headings.append(txt)
        else:
            lines.append(txt)

    # 表格内容以管道分隔追加，避免丢失
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    text = "\n".join(lines).strip()
    if not text:
        raise ParseError("DOCX 未提取到文本内容")
    return ParsedDocument(text=text, headings=headings)


def parse(path: Path, ext: str) -> ParsedDocument:
    """按扩展名解析文档。"""
    ext = ext.lower().lstrip(".")
    if ext == "pdf":
        return _parse_pdf(path)
    if ext == "docx":
        return _parse_docx(path)
    if ext == "svg":
        # SVG 是 XML 文本，直接按文本读，不必 OCR
        return _parse_text(path)
    if ext in files.OCR_IMAGE_EXTENSIONS:
        return _parse_image(path)
    if ext in TEXT_EXTENSIONS:
        return _parse_text(path)
    raise ParseError(f"暂不支持解析的格式：.{ext}")
